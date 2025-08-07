from excel.register_form_page import register
import openpyxl

file = openpyxl.load_workbook("Playwright/excel/users.xlsx")

def rows(sheet_name):
    rows_number = file[sheet_name]
    return rows_number.max_row

def values(sheet_name, row, column):
    rows_number = file[sheet_name]
    column = rows_number.cell(int(row), int(column))
    return column.value

def test_fill_out_form(set_up_excel) -> None:
    page = set_up_excel
    ex = register(page)

    for n in range(3, 6):
        name = values("users", n, 1)
        last_name = values("users", n, 2)
        phone = values("users", n, 3)
        country = values("users", n, 4)
        email = values("users", n, 5)
        password = values("users", n, 6)
        
        print(f"User {name, last_name, phone, country, email, password}")

        ex.enter_name(name)
        ex.enter_last_name(last_name)
        ex.enter_phone_number(str(phone))
        ex.select_country(country)
        ex.enter_email(email)
        ex.enter_password(password)
        ex.check_terms_box()
        ex.submit_register()
        ex.read_confirmation_msg("The account has been successfully created!")
        page.screenshot(path=f"Playwright/excel/screenshots/register{n}.png")
